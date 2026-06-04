"""
Parsers for binary document formats (Office, PDF).
Each parser accepts raw bytes and returns extracted text.
Returns None for unrecognised extensions so the caller can fall back.
"""
from __future__ import annotations

import io


def extract_document_text(filename: str, data: bytes) -> str | None:
    """Dispatch to the appropriate parser based on file extension.

    Returns extracted text, or None if the extension is not supported.
    """
    ext = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    if ext == "pdf":
        return _parse_pdf(data)
    if ext in ("xlsx", "xlsm", "xltx", "xltm"):
        return _parse_xlsx(data)
    if ext == "xls":
        return _parse_xls(data)
    if ext in ("docx", "docm"):
        return _parse_docx(data)
    if ext in ("pptx", "pptm"):
        return _parse_pptx(data)
    return None


# ---------------------------------------------------------------------------
# PDF
# ---------------------------------------------------------------------------

_MAX_PDF_PAGES = 500
_MAX_SHEET_ROWS = 10_000


def _parse_pdf(data: bytes) -> str:
    import pymupdf  # type: ignore

    parts: list[str] = []
    with pymupdf.open(stream=data, filetype="pdf") as doc:
        page_count = doc.page_count
        limit = min(page_count, _MAX_PDF_PAGES)
        for i in range(limit):
            text = doc[i].get_text().strip()
            if text:
                parts.append(f"=== Page {i + 1} ===\n{text}")
        if page_count > _MAX_PDF_PAGES:
            parts.append(
                f"[Truncated — extracted {_MAX_PDF_PAGES} of {page_count} pages]"
            )
    return "\n\n".join(parts) if parts else "[PDF contains no extractable text]"


# ---------------------------------------------------------------------------
# Excel
# ---------------------------------------------------------------------------

def _parse_xlsx(data: bytes) -> str:
    import openpyxl  # type: ignore

    wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    parts: list[str] = []
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        rows: list[str] = []
        truncated = False
        for row in ws.iter_rows(values_only=True):
            if len(rows) >= _MAX_SHEET_ROWS:
                truncated = True
                break
            if any(cell is not None for cell in row):
                rows.append("\t".join("" if cell is None else str(cell) for cell in row))
        if rows:
            section = f"=== Sheet: {sheet_name} ===\n" + "\n".join(rows)
            if truncated:
                section += f"\n[Truncated — showing first {_MAX_SHEET_ROWS} rows]"
            parts.append(section)
    wb.close()
    return "\n\n".join(parts) if parts else "[Workbook contains no data]"


def _parse_xls(data: bytes) -> str:
    import xlrd  # type: ignore

    wb = xlrd.open_workbook(file_contents=data)
    parts: list[str] = []
    for sheet_name in wb.sheet_names():
        ws = wb.sheet_by_name(sheet_name)
        rows: list[str] = []
        row_limit = min(ws.nrows, _MAX_SHEET_ROWS)
        for row_idx in range(row_limit):
            cells = [str(ws.cell_value(row_idx, col)) for col in range(ws.ncols)]
            if any(c.strip() for c in cells):
                rows.append("\t".join(cells))
        if rows:
            section = f"=== Sheet: {sheet_name} ===\n" + "\n".join(rows)
            if ws.nrows > _MAX_SHEET_ROWS:
                section += f"\n[Truncated — showing first {_MAX_SHEET_ROWS} rows]"
            parts.append(section)
    return "\n\n".join(parts) if parts else "[Workbook contains no data]"


# ---------------------------------------------------------------------------
# Word
# ---------------------------------------------------------------------------

def _parse_docx(data: bytes) -> str:
    import docx  # type: ignore  # python-docx

    doc = docx.Document(io.BytesIO(data))
    lines = [para.text for para in doc.paragraphs if para.text.strip()]
    return "\n".join(lines) if lines else "[Document contains no text]"


# ---------------------------------------------------------------------------
# PowerPoint
# ---------------------------------------------------------------------------

def _parse_pptx(data: bytes) -> str:
    from pptx import Presentation  # type: ignore  # python-pptx

    prs = Presentation(io.BytesIO(data))
    parts: list[str] = []
    for i, slide in enumerate(prs.slides, start=1):
        texts: list[str] = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                for para in shape.text_frame.paragraphs:
                    text = para.text.strip()
                    if text:
                        texts.append(text)
        if texts:
            parts.append(f"=== Slide {i} ===\n" + "\n".join(texts))
    return "\n\n".join(parts) if parts else "[Presentation contains no text]"
